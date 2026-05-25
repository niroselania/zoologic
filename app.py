from __future__ import annotations

import os
import shutil
import unicodedata
import json
from io import BytesIO
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook
from werkzeug.utils import secure_filename


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("DATA_DIR", BASE_DIR / "data"))
UPLOAD_DIR = DATA_DIR / "uploads"
BACKUP_DIR = DATA_DIR / "backups"
CURRENT_FILE = DATA_DIR / "zoologic_actual.xlsx"
HISTORY_FILE = DATA_DIR / "historial_cambios.jsonl"
ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 30 * 1024 * 1024


def ensure_dirs() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.casefold().split())


def serialize_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return value


def workbook_path() -> Path:
    if not CURRENT_FILE.exists():
        raise FileNotFoundError("Todavia no hay una planilla cargada.")
    return CURRENT_FILE


def open_workbook():
    return load_workbook(workbook_path())


def header_score(values: list[Any]) -> int:
    normalized = [normalize_text(value) for value in values if value not in (None, "")]
    text = " | ".join(normalized)
    score = len(normalized)
    for keyword in ("usuario", "puesto", "serie", "concepto", "canal", "estado"):
        if keyword in text:
            score += 8
    if "usuario" in text and "puesto" in text:
        score += 30
    return score


def detect_header_row(sheet) -> int:
    best_row = 1
    best_score = -1
    limit = min(sheet.max_row, 15)
    for row_number in range(1, limit + 1):
        values = [sheet.cell(row_number, col).value for col in range(1, sheet.max_column + 1)]
        score = header_score(values)
        if score > best_score:
            best_row = row_number
            best_score = score
    return best_row


def read_headers(sheet, header_row: int) -> list[dict[str, Any]]:
    headers: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    for column in range(1, sheet.max_column + 1):
        raw_value = sheet.cell(header_row, column).value
        if raw_value in (None, ""):
            continue
        label = str(raw_value).strip()
        count = seen.get(label, 0) + 1
        seen[label] = count
        key = label if count == 1 else f"{label} ({count})"
        headers.append({"key": key, "label": label, "column": column})
    return headers


def searchable_headers(headers: list[dict[str, Any]]) -> list[dict[str, Any]]:
    matches = []
    for header in headers:
        text = normalize_text(header["label"])
        is_user = "usuario" in text and not text.startswith("fecha")
        is_position = "puesto" in text
        if is_user or is_position:
            matches.append(header)
    return matches


def header_matches(header: dict[str, Any], *terms: str) -> bool:
    label = normalize_text(header["label"])
    return all(term in label for term in terms)


def find_header(headers: list[dict[str, Any]], *terms: str) -> dict[str, Any] | None:
    for header in headers:
        if header_matches(header, *terms):
            return header
    return None


def value_for(headers: list[dict[str, Any]], values: dict[str, Any], *terms: str) -> Any:
    header = find_header(headers, *terms)
    return values.get(header["key"]) if header else None


def row_remote_info(headers: list[dict[str, Any]], values: dict[str, Any]) -> dict[str, Any]:
    remote_header = find_header(headers, "escritorio", "remoto")
    password_header = find_header(headers, "password") or find_header(headers, "contrasena")
    remote_value = values.get(remote_header["key"]) if remote_header else None
    password_value = values.get(password_header["key"]) if password_header else None
    return {
        "available": bool(str(remote_value or "").strip()),
        "address": remote_value,
        "password": password_value,
    }


def row_access_summary(headers: list[dict[str, Any]], values: dict[str, Any]) -> dict[str, Any]:
    return {
        "puesto": value_for(headers, values, "puesto"),
        "usuario": value_for(headers, values, "usuario"),
        "serie": value_for(headers, values, "serie"),
        "canal": value_for(headers, values, "canal"),
        "estado": value_for(headers, values, "estado"),
        "sitio": value_for(headers, values, "sitio") or value_for(headers, values, "ubicacion"),
        "escritorio_remoto": value_for(headers, values, "escritorio", "remoto"),
        "password": value_for(headers, values, "password") or value_for(headers, values, "contrasena"),
    }


def row_issues(
    headers: list[dict[str, Any]],
    values: dict[str, Any],
    duplicate_series: set[str],
) -> list[str]:
    issues = []
    usuario = value_for(headers, values, "usuario")
    puesto = value_for(headers, values, "puesto")
    serie = value_for(headers, values, "serie")
    estado = value_for(headers, values, "estado")
    remote = value_for(headers, values, "escritorio", "remoto")
    password = value_for(headers, values, "password") or value_for(headers, values, "contrasena")

    if find_header(headers, "usuario") and not str(usuario or "").strip():
        issues.append("Usuario vacio")
    if find_header(headers, "puesto") and not str(puesto or "").strip():
        issues.append("Puesto vacio")
    if find_header(headers, "escritorio", "remoto") and not str(remote or "").strip():
        issues.append("Escritorio remoto vacio")
    if (find_header(headers, "password") or find_header(headers, "contrasena")) and not str(password or "").strip():
        issues.append("Password vacio")
    if str(serie or "").strip() and str(serie).strip() in duplicate_series:
        issues.append("Serie duplicada")
    if str(estado or "").strip() and normalize_text(estado) not in {"activado", "activo"}:
        issues.append(f"Estado: {estado}")

    return issues


def build_dashboard(
    headers: list[dict[str, Any]],
    all_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    series: dict[str, int] = {}
    canales: dict[str, int] = {}
    estados: dict[str, int] = {}
    sitios: dict[str, int] = {}

    for row in all_rows:
        values = row["values"]
        serie = str(value_for(headers, values, "serie") or "").strip()
        canal = str(value_for(headers, values, "canal") or "").strip()
        estado = str(value_for(headers, values, "estado") or "").strip()
        sitio = str(value_for(headers, values, "sitio") or value_for(headers, values, "ubicacion") or "").strip()
        if serie:
            series[serie] = series.get(serie, 0) + 1
        if canal:
            canales[canal] = canales.get(canal, 0) + 1
        if estado:
            estados[estado] = estados.get(estado, 0) + 1
        if sitio:
            sitios[sitio] = sitios.get(sitio, 0) + 1

    duplicate_series = {serie for serie, count in series.items() if count > 1}
    review_rows = []
    issue_totals: dict[str, int] = {}
    for row in all_rows:
        issues = row_issues(headers, row["values"], duplicate_series)
        for issue in issues:
            issue_totals[issue] = issue_totals.get(issue, 0) + 1
        if issues:
            summary = row_access_summary(headers, row["values"])
            review_rows.append(
                {
                    "row_id": row["row_id"],
                    "issues": issues,
                    "puesto": summary["puesto"],
                    "usuario": summary["usuario"],
                    "serie": summary["serie"],
                    "estado": summary["estado"],
                }
            )

    def top_items(items: dict[str, int], limit: int = 8) -> list[dict[str, Any]]:
        return [
            {"name": name, "count": count}
            for name, count in sorted(items.items(), key=lambda item: (-item[1], item[0]))[:limit]
        ]

    return {
        "total": len(all_rows),
        "active": sum(1 for row in all_rows if normalize_text(value_for(headers, row["values"], "estado")) in {"activado", "activo"}),
        "missing_user": issue_totals.get("Usuario vacio", 0),
        "missing_remote": issue_totals.get("Escritorio remoto vacio", 0),
        "missing_password": issue_totals.get("Password vacio", 0),
        "duplicate_series": len(duplicate_series),
        "issue_count": sum(issue_totals.values()),
        "issue_totals": top_items(issue_totals, 10),
        "channels": top_items(canales),
        "states": top_items(estados),
        "sites": top_items(sitios),
        "review_rows": review_rows[:30],
    }


def preferred_sheet_name(workbook) -> str:
    for sheet in workbook.worksheets:
        header_row = detect_header_row(sheet)
        headers = read_headers(sheet, header_row)
        labels = " | ".join(normalize_text(header["label"]) for header in headers)
        if "usuario" in labels and "puesto" in labels:
            return sheet.title
    return workbook.active.title


def sheet_names(workbook) -> list[str]:
    return [sheet.title for sheet in workbook.worksheets]


def rows_from_sheet(sheet_name: str | None = None, search: str = "", search_mode: str = "targeted") -> dict[str, Any]:
    workbook = open_workbook()
    try:
        selected_sheet = sheet_name if sheet_name in workbook.sheetnames else preferred_sheet_name(workbook)
        sheet = workbook[selected_sheet]
        header_row = detect_header_row(sheet)
        headers = read_headers(sheet, header_row)
        search_headers = searchable_headers(headers)
        search_key = normalize_text(search)
        all_rows = []
        rows = []

        for row_number in range(header_row + 1, sheet.max_row + 1):
            values = {}
            has_value = False
            for header in headers:
                value = serialize_value(sheet.cell(row_number, header["column"]).value)
                values[header["key"]] = value
                if value not in (None, ""):
                    has_value = True

            if not has_value:
                continue

            row_payload = {
                "row_id": row_number,
                "values": values,
                "remote_desktop": row_remote_info(headers, values),
                "access_summary": row_access_summary(headers, values),
            }
            all_rows.append(row_payload)

            if search_key:
                haystack_headers = headers if search_mode == "global" else search_headers or headers
                haystack = " | ".join(
                    normalize_text(sheet.cell(row_number, header["column"]).value)
                    for header in haystack_headers
                )
                if search_key not in haystack:
                    continue

            rows.append(row_payload)

        return {
            "filename": CURRENT_FILE.name,
            "sheet": sheet.title,
            "sheets": sheet_names(workbook),
            "header_row": header_row,
            "headers": headers,
            "search_columns": [header["label"] for header in search_headers],
            "search_mode": search_mode,
            "dashboard": build_dashboard(headers, all_rows),
            "rows": rows,
            "total_rows": len(all_rows),
            "filtered_rows": len(rows),
        }
    finally:
        workbook.close()


def make_backup() -> None:
    if CURRENT_FILE.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(CURRENT_FILE, BACKUP_DIR / f"zoologic_backup_{stamp}.xlsx")


def append_history(sheet_name: str, row_id: int, changes: list[dict[str, Any]]) -> None:
    if not changes:
        return
    ensure_dirs()
    entry = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "sheet": sheet_name,
        "row_id": row_id,
        "changes": changes,
    }
    with HISTORY_FILE.open("a", encoding="utf-8") as file:
        file.write(json.dumps(entry, ensure_ascii=False) + "\n")


def read_history(limit: int = 80) -> list[dict[str, Any]]:
    if not HISTORY_FILE.exists():
        return []
    lines = HISTORY_FILE.read_text(encoding="utf-8").splitlines()
    entries = []
    for line in lines[-limit:]:
        try:
            entries.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return list(reversed(entries))


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/status")
def status():
    if not CURRENT_FILE.exists():
        return jsonify({"loaded": False})
    data = rows_from_sheet()
    return jsonify(
        {
            "loaded": True,
            "filename": data["filename"],
            "sheet": data["sheet"],
            "sheets": data["sheets"],
            "headers": data["headers"],
            "search_columns": data["search_columns"],
            "total_rows": data["total_rows"],
        }
    )


@app.post("/api/upload")
def upload_file():
    ensure_dirs()
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "Subi un archivo Excel para continuar."}), 400

    suffix = Path(file.filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        return jsonify({"error": "Formato no soportado. Usa .xlsx o .xlsm."}), 400

    filename = secure_filename(file.filename)
    uploaded_path = UPLOAD_DIR / filename
    file.save(uploaded_path)

    try:
        workbook = load_workbook(uploaded_path)
        workbook.close()
    except Exception as exc:
        uploaded_path.unlink(missing_ok=True)
        return jsonify({"error": f"No pude abrir el Excel: {exc}"}), 400

    make_backup()
    shutil.copy2(uploaded_path, CURRENT_FILE)
    return jsonify(rows_from_sheet())


@app.get("/api/rows")
def rows():
    try:
        return jsonify(rows_from_sheet(request.args.get("sheet"), request.args.get("q", ""), request.args.get("mode", "targeted")))
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404


@app.get("/api/history")
def history():
    return jsonify({"history": read_history()})


@app.patch("/api/rows/<int:row_id>")
def update_row(row_id: int):
    payload = request.get_json(silent=True) or {}
    values = payload.get("values", {})
    sheet_name = payload.get("sheet")
    if not isinstance(values, dict):
        return jsonify({"error": "Los datos enviados no son validos."}), 400

    try:
        workbook = open_workbook()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    try:
        selected_sheet = sheet_name if sheet_name in workbook.sheetnames else preferred_sheet_name(workbook)
        sheet = workbook[selected_sheet]
        header_row = detect_header_row(sheet)
        headers = read_headers(sheet, header_row)
        key_to_column = {header["key"]: header["column"] for header in headers}

        if row_id <= header_row or row_id > sheet.max_row:
            return jsonify({"error": "La fila indicada no existe."}), 404

        unknown = [key for key in values if key not in key_to_column]
        if unknown:
            return jsonify({"error": f"Columnas desconocidas: {', '.join(unknown)}"}), 400

        changes = []
        for key, value in values.items():
            old_value = serialize_value(sheet.cell(row=row_id, column=key_to_column[key]).value)
            new_value = value if value != "" else None
            if str(old_value or "") != str(new_value or ""):
                changes.append({"field": key, "old": old_value, "new": new_value})

        if not changes:
            return jsonify({"ok": True, "changed": False})

        make_backup()
        for key, value in values.items():
            sheet.cell(row=row_id, column=key_to_column[key]).value = value if value != "" else None

        workbook.save(CURRENT_FILE)
        append_history(selected_sheet, row_id, changes)
        return jsonify({"ok": True, "changed": True})
    finally:
        workbook.close()


@app.get("/api/rdp/<int:row_id>")
def rdp_file(row_id: int):
    sheet_name = request.args.get("sheet")

    try:
        workbook = open_workbook()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404

    try:
        selected_sheet = sheet_name if sheet_name in workbook.sheetnames else preferred_sheet_name(workbook)
        sheet = workbook[selected_sheet]
        header_row = detect_header_row(sheet)
        headers = read_headers(sheet, header_row)

        if row_id <= header_row or row_id > sheet.max_row:
            return jsonify({"error": "La fila indicada no existe."}), 404

        values = {
            header["key"]: serialize_value(sheet.cell(row_id, header["column"]).value)
            for header in headers
        }
        remote_info = row_remote_info(headers, values)
        address = str(remote_info["address"] or "").strip()
        if not address:
            return jsonify({"error": "Esta fila no tiene dato en Escritorio remoto."}), 404

        username_header = find_header(headers, "usuario") or find_header(headers, "persona")
        username = str(values.get(username_header["key"]) or "").strip() if username_header else ""
        rdp_lines = [
            "screen mode id:i:2",
            "use multimon:i:0",
            "desktopwidth:i:1920",
            "desktopheight:i:1080",
            "session bpp:i:32",
            "prompt for credentials:i:1",
            "authentication level:i:2",
            f"full address:s:{address}",
        ]
        if username:
            rdp_lines.append(f"username:s:{username}")

        payload = "\r\n".join(rdp_lines) + "\r\n"
        stream = BytesIO(payload.encode("utf-16le"))
        safe_name = "".join(char for char in address if char.isalnum() or char in ("-", "_", ".")).strip(".")
        filename = f"zoologic_{safe_name or row_id}.rdp"
        return send_file(stream, mimetype="application/rdp", as_attachment=True, download_name=filename)
    finally:
        workbook.close()


@app.get("/api/download")
def download():
    try:
        path = workbook_path()
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 404
    return send_file(path, as_attachment=True, download_name="zoologic_actualizado.xlsx")


if __name__ == "__main__":
    ensure_dirs()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8000")), debug=False)
